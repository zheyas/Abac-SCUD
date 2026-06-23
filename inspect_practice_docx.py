from docx import Document


path = "/Users/evgenijasakov/Desktop/учеба/Практика/примеры/Рецензия с подписями.docx"
document = Document(path)

print(
    "PARAGRAPHS", len(document.paragraphs),
    "TABLES", len(document.tables),
    "SECTIONS", len(document.sections),
)
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if text:
        print(f"P{index}: {text}")
        for run_index, run in enumerate(paragraph.runs):
            if run.text:
                print(
                    f"  RUN{run_index}: {run.text!r} "
                    f"bold={run.bold} underline={run.underline} italic={run.italic} "
                    f"size={run.font.size}"
                )

for table_index, table in enumerate(document.tables):
    print(f"\nTABLE {table_index} rows={len(table.rows)} cols={len(table.columns)}")
    for row_index, row in enumerate(table.rows):
        values = [" ".join(cell.text.split()) for cell in row.cells]
        print(f"R{row_index}: " + " || ".join(values))
