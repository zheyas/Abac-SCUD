from copy import deepcopy
from pathlib import Path

from docx import Document


SOURCE = Path("/Users/evgenijasakov/НИР/Дневник_и_задание_Яковлев_2026.docx")
TEMPLATE = Path("/Users/evgenijasakov/НИР/Дневник_Задание_рабочая_копия.docx")
OUTPUT = Path("/Users/evgenijasakov/НИР/Дневник_и_задание_Ясаков_2026.docx")


def set_full(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def replace_in_runs(paragraph):
    replacements = {
        "Яковлева": "Ясакова",
        "Яковлев": "Ясаков",
    }
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, new)


document = Document(SOURCE)
template = Document(TEMPLATE)

# Restore the complete review block from the untouched source template.
current_paragraphs = document.paragraphs
template_paragraphs = template.paragraphs
for index in range(101, 139):
    current_element = current_paragraphs[index]._p
    current_element.getparent().replace(
        current_element,
        deepcopy(template_paragraphs[index]._p),
    )

# Correct the surname everywhere while preserving the user's run formatting.
for paragraph in document.paragraphs:
    replace_in_runs(paragraph)
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_runs(paragraph)

p = document.paragraphs
set_full(p[123], "Студент Ясаков Евгений Михайлович")
set_full(p[138], "«______» ______________ 2026 г.")

document.core_properties.author = "Ясаков Евгений Михайлович"
document.save(OUTPUT)
print(OUTPUT)
