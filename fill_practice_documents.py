from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt


SOURCE = Path("/Users/evgenijasakov/НИР/Дневник_Задание_рабочая_копия.docx")
OUTPUT = Path("/Users/evgenijasakov/НИР/Дневник_и_задание_Яковлев_2026.docx")


def set_full(paragraph, text):
    """Replace visible text while retaining the first run's formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def set_label_value(paragraph, label, value, underline=True):
    """Keep a form label and replace its underlined value."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(label)
        value_run = paragraph.add_run(value)
    else:
        runs[0].text = label
        if len(runs) > 1:
            value_run = runs[1]
            value_run.text = value
            for run in runs[2:]:
                run.text = ""
        else:
            value_run = paragraph.add_run(value)
    value_run.underline = underline


def fill_signature_name(paragraph, name):
    """Fill the last underlined space reserved for initials and surname."""
    candidates = [
        run for run in paragraph.runs
        if run.underline and run.text and not run.text.strip() and "\t" not in run.text
    ]
    if candidates:
        candidates[-1].text = name
        return
    run = paragraph.add_run(name)
    run.underline = True


def set_cell(cell, text):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    for extra_paragraph in cell.paragraphs[1:]:
        set_full(extra_paragraph, "")


document = Document(SOURCE)
p = document.paragraphs

# Individual assignment.
set_full(
    p[0],
    "Федеральное государственное бюджетное образовательное учреждение  "
    "\nвысшего образования «Оренбургский государственный университет имени В.А. Бондаренко» (ОГУ)",
)
p[0].runs[0].font.size = Pt(10)
set_label_value(
    p[3],
    "Вид, тип практики ",
    "Производственная практика (научно-исследовательская работа)",
)
set_label_value(p[4], "Обучающийся ", "Яковлев Евгений Михайлович")
set_label_value(p[6], "Курс ", "4")
set_label_value(p[7], "Институт ", "Математики и информационных технологий")
set_label_value(p[8], "Форма обучения ", "Очная")
set_label_value(p[9], "Специальность или направление подготовки ", "10.05.01 Компьютерная безопасность")
set_label_value(p[10], "Направленность (профиль) ", "Разработка защищенного программного обеспечения")
set_label_value(
    p[11],
    "Тема научно-исследовательской работы ",
    "Разработка прототипа интерактивной системы контроля и управления доступом для промышленного объекта",
)
set_label_value(
    p[12],
    "Цель научно-исследовательской работы ",
    "Разработать и протестировать программный прототип СКУД для промышленного объекта",
)

assignment_items = [
    "Проанализировать требования к программному прототипу СКУД",
    "Спроектировать модель пользователей, групп, зданий и правил доступа",
    "Реализовать серверную часть и механизмы разграничения доступа",
    "Разработать интерактивную 3D-карту объектов",
    "Провести тестирование и подготовить документацию",
]
for index, text in enumerate(assignment_items, start=14):
    set_full(p[index], text)

set_label_value(p[19], "Дата выдачи задания ", "22 июня 2026")
fill_signature_name(p[23], "А. А. Рязанов")
fill_signature_name(p[26], "Яковлев Е. М.")
set_full(
    p[29],
    "Индивидуальное задание выполнено в полном объёме; программный прототип разработан и протестирован.",
)
set_full(p[30], "")
fill_signature_name(p[34], "А. А. Рязанов")
if len(p[35].runs) > 4:
    p[35].runs[4].text = ""

# Diary title page.
set_full(p[39], "«Оренбургский государственный университет имени В.А. Бондаренко»")
set_label_value(p[53], "Студента ", "Яковлева Евгения Михайловича")
set_label_value(p[56], "группы ", "22КБ(с)РЗПО-2")
set_full(p[63], "Филиал «Оренбургский гелиевый завод» ООО «Газпром переработка»")
set_label_value(
    p[66],
    "Срок практики с ",
    "«22» июня 2026 г. по «04» июля 2026 г.",
)
set_full(p[81], "Оренбург 2026")

# Calendar plan and completed-work journal.
plan = [
    (
        "22.06.2026",
        "Ознакомление с программой и задачами практики. Инструктаж по охране труда и режиму конфиденциальности.",
        "Анализ требований к программному прототипу СКУД.",
    ),
    (
        "23.06.2026 - 24.06.2026",
        "Исследование подходов к разграничению доступа. Проектирование структуры данных.",
        "Определены сущности пользователей, групп, зданий и правил доступа.",
    ),
    (
        "25.06.2026 - 26.06.2026",
        "Разработка серверной части и механизма расчёта групповых и персональных прав доступа.",
        "Реализованы API, режимы зданий и журнал событий.",
    ),
    (
        "29.06.2026 - 01.07.2026",
        "Разработка интерактивной 3D-карты и отображения состояния объектов в реальном времени.",
        "Реализованы индикация доступности, моделирование времени и погодных условий.",
    ),
    (
        "02.07.2026 - 04.07.2026",
        "Функциональное тестирование, устранение замечаний и подготовка документации.",
        "Проверены сценарии допуска и отказа; подготовлены дневник и отчёт.",
    ),
]
for row_index, (date, work, note) in enumerate(plan, start=1):
    set_cell(document.tables[0].cell(row_index, 0), str(row_index))
    set_cell(document.tables[0].cell(row_index, 1), date)
    set_cell(document.tables[0].cell(row_index, 2), work)
    set_cell(document.tables[0].cell(row_index, 3), note)

completed = [
    "Проведено ознакомление с программой и задачами практики, требованиями охраны труда и режимом конфиденциальности.",
    "Исследованы подходы к разграничению доступа; спроектирована структура данных программного прототипа.",
    "Разработаны серверная часть, механизм групповых и персональных прав доступа, режимы зданий и журнал событий.",
    "Разработана интерактивная 3D-карта с отображением доступности объектов, времени суток и погодных условий.",
    "Выполнено функциональное тестирование, устранены замечания и подготовлена документация по проекту.",
]
for row_index, ((date, _, _), work) in enumerate(zip(plan, completed), start=1):
    set_cell(document.tables[1].cell(row_index, 0), str(row_index))
    set_cell(document.tables[1].cell(row_index, 1), date)
    set_cell(document.tables[1].cell(row_index, 2), work)
    set_cell(document.tables[1].cell(row_index, 3), "")

fill_signature_name(p[87], "А. А. Рязанов")
set_full(p[89], "«04» июля 2026 г.")
fill_signature_name(p[97], "А. А. Рязанов")
set_full(p[100], "«04» июля 2026 г.")

# Supervisor review.
set_full(p[105], "В ходе производственной практики была поставлена и достигнута цель:")
set_full(p[107], "разработать и апробировать программный прототип интерактивной системы контроля")
set_full(p[109], "и управления доступом для промышленного объекта, обеспечивающий настройку")
set_full(p[111], "групповых и персональных прав, режимов зданий и учёт событий прохода.")
set_full(p[113], "В процессе достижения цели получены следующие навыки:")
set_full(p[115], "анализ требований и проектирование модели разграничения доступа;")
set_full(p[117], "разработка серверной части на Python/Flask и базы данных SQLite;")
set_full(p[119], "создание 3D-интерфейса, тестирование и документирование программного решения.")
set_full(p[123], "Студент Яковлев Евгений Михайлович")
set_label_value(
    p[128],
    "подготовке от профильной организации ",
    "_________            А. А. Рязанов",
)
set_full(p[138], "«04» июля 2026 г.")

document.core_properties.title = "Дневник и индивидуальное задание на производственную практику"
document.core_properties.subject = "Практика в филиале «Оренбургский гелиевый завод» ООО «Газпром переработка»"
document.core_properties.author = "Яковлев Евгений Михайлович"
document.save(OUTPUT)
print(OUTPUT)
