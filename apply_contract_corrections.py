from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path("/Users/evgenijasakov/НИР/Дневник_и_задание_Ясаков_2026.docx")
OUTPUT = Path("/Users/evgenijasakov/НИР/Комплект_практической_подготовки_Ясаков_2026.docx")


def set_full(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def set_label_value(paragraph, label, value, value_underlined=True):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(label)
        value_run = paragraph.add_run(value)
    else:
        runs[0].text = label
        if len(runs) > 1:
            value_run = runs[1]
            value_run.text = value
            for run in runs[2:]:
                run.text = ""
        else:
            value_run = paragraph.add_run(value)
    value_run.underline = value_underlined


def set_signature_line(paragraph, label, name="", font_size=None):
    set_full(paragraph, label)
    line_run = paragraph.add_run(" ____________________ ")
    line_run.underline = False
    if name:
        name_run = paragraph.add_run(name)
        name_run.underline = True
    if font_size:
        for run in paragraph.runs:
            run.font.size = Pt(font_size)


def set_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


document = Document(SOURCE)
p = document.paragraphs

# Terminology required by the practical-training agreement and supervisor's notes.
set_full(p[2], "ИНДИВИДУАЛЬНОЕ ЗАДАНИЕ НА ПРАКТИЧЕСКУЮ ПОДГОТОВКУ")
set_label_value(
    p[3],
    "Вид, тип практической подготовки ",
    "Производственная практика (научно-исследовательская работа)",
)
set_full(
    p[13],
    "Содержание задания на практическую подготовку (перечень подлежащих рассмотрению вопросов, "
    "выполняемых работ, связанных с будущей профессиональной деятельностью):",
)
set_full(
    p[28],
    "Заключение руководителя по практической подготовке о выполнении задания практической подготовки:",
)

# The supervisor explicitly asked not to identify the responsible employee in the assignment.
for index in (23, 24, 34, 35):
    set_full(p[index], "")

# Diary and report terminology. The component name remains exactly as in the agreement.
set_full(p[44], "ДНЕВНИК ПРАКТИЧЕСКОЙ ПОДГОТОВКИ")
set_full(p[45], "(ПРОИЗВОДСТВЕННАЯ ПРАКТИКА, НАУЧНО-ИССЛЕДОВАТЕЛЬСКАЯ РАБОТА)")
set_full(p[62], "Место практической подготовки")
set_full(
    p[63],
    "Филиал «Оренбургский гелиевый завод» ООО «Газпром переработка», "
    "Бюро алгоритмизации и программирования",
)
set_label_value(
    p[66],
    "Срок практической подготовки: ",
    "«22» июня - «04» июля 2026 г.",
)
for index in (63, 66):
    for run in p[index].runs:
        run.font.size = Pt(12)

# Keep the city and year on the diary title page, then start the plan on a new page.
for index in (78, 79, 80):
    p[index].paragraph_format.line_spacing = Pt(1)
    p[index].paragraph_format.space_before = Pt(0)
    p[index].paragraph_format.space_after = Pt(0)
    for run in p[index].runs:
        run.font.size = Pt(1)
p[82].paragraph_format.page_break_before = True
set_full(p[82], "Календарный план практической подготовки")
set_full(p[92], "Отчёт о практической подготовке")
set_full(p[101], "Отзыв ответственного за практическую подготовку")
p[101].paragraph_format.page_break_before = True
set_full(p[103], "Общая характеристика выполненной работы")
set_full(p[105], "В ходе практической подготовки была поставлена и достигнута цель")
set_full(p[121], "Заключение ответственного")

# Signature blocks: the individual assignment stays blank, while plan/report identify A. A. Ryazanov.
set_signature_line(p[85], "Руководитель по практической подготовке от ОГУ", "И.В. Влацкая", 10.5)
set_signature_line(
    p[87],
    "Ответственный за практическую подготовку от профильной организации",
    "А. А. Рязанов",
    9.5,
)
set_signature_line(p[95], "Руководитель по практической подготовке от ОГУ", "И.В. Влацкая", 10.5)
set_signature_line(
    p[97],
    "Ответственный за практическую подготовку от профильной организации",
    "А. А. Рязанов",
    9.5,
)
set_full(p[127], "Ответственный за практическую")
set_label_value(
    p[128],
    "подготовку от профильной организации ",
    "_________            А. А. Рязанов",
)

# Use the agreement's terminology in the plan and report tables as well.
for table_index in (0, 1):
    for paragraph in document.tables[table_index].cell(1, 2).paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace("задачами практики", "задачами практической подготовки")

# Add a separate unsigned review page based on the supplied example.
document.add_page_break()

title_lines = [
    "РЕЦЕНЗИЯ",
    "на результаты практической подготовки",
    "(производственная практика, научно-исследовательская работа)",
    "обучающегося Ясакова Евгения Михайловича",
    "на тему «Разработка прототипа интерактивной системы контроля и управления доступом "
    "для промышленного объекта»",
]
for index, text in enumerate(title_lines):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4 if index else 6)
    run = paragraph.add_run(text)
    set_font(run, 14, True)

review_paragraphs = [
    "В период практической подготовки Ясаков Евгений Михайлович выполнил "
    "научно-исследовательскую работу и разработал программный прототип интерактивной "
    "системы контроля и управления доступом для промышленного объекта.",
    "В работе проанализированы требования к разграничению доступа, спроектирована модель "
    "пользователей, групп, зданий, правил доступа и рабочих смен. Реализованы серверная часть, "
    "групповые и персональные права доступа, режимы работы зданий и журнал событий.",
    "Разработана интерактивная 3D-карта территории с отображением доступности объектов в "
    "реальном времени. Проведено функциональное тестирование сценариев допуска и отказа, "
    "устранены выявленные замечания и подготовлена документация по проекту.",
    "Недостатков, препятствующих положительной оценке научно-исследовательской работы, не выявлено. "
    "Работа по содержанию, объёму и степени проработки соответствует установленным требованиям. "
    "Поставленные цель и задачи достигнуты, полученные результаты обоснованы.",
]
for text in review_paragraphs:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_font(run, 14)

paragraph = document.add_paragraph()
paragraph.paragraph_format.space_before = Pt(6)
run = paragraph.add_run("Обучающийся Ясаков Е. М. заслуживает оценку ____________________")
set_font(run, 14)

paragraph = document.add_paragraph()
run = paragraph.add_run("Рецензент: руководитель по практической подготовке от ОГУ,")
set_font(run, 14)
paragraph = document.add_paragraph()
run = paragraph.add_run(
    "заведующий кафедрой компьютерной безопасности и математического обеспечения "
    "информационных систем И. В. Влацкая ____________________"
)
set_font(run, 14)

paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
paragraph.paragraph_format.space_before = Pt(10)
run = paragraph.add_run("«______» ______________ 2026 г.")
set_font(run, 14)

document.core_properties.title = "Комплект документов по практической подготовке"
document.core_properties.subject = (
    "Практическая подготовка в Филиале «Оренбургский гелиевый завод» "
    "ООО «Газпром переработка», Бюро алгоритмизации и программирования"
)
document.core_properties.author = "Ясаков Евгений Михайлович"
document.save(OUTPUT)
print(OUTPUT)
